#!/usr/bin/env python3
"""Remplit le questionnaire 'Cohort 1 Innovator Story' pour Forgeron."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

SRC = Path("/tmp/claude-0/-home-user-forgeron/f115b2cf-c02c-5ae4-8f73-96461eacf843/scratchpad/"
           "q_in/e2dd4a11-METI_AI_Questionnaire.docx")
OUT = Path("/home/user/forgeron/scratch/METI_AI_Questionnaire_Forgeron_LamineSACKO.docx")

Q1 = "Manufacturing and industrial automation — advanced manufacturing technologies."

Q2 = ("Lamine SACKO. Mali. UniPod Mali, hosted by the École Nationale d'Ingénieurs "
      "Abderhamane Baba Touré (ENI-ABT) in Bamako.")

Q3 = "As a team, of two."

TEAM = [
    ("Lamine SACKO",
     "Founder and engineer. Designed the machine and built the control software, the mobile "
     "application and the AI agent."),
    ("Aboubacar Diamouténé",
     "Electronics and assembly technician. Builds and wires the machine, and diagnoses the "
     "faults that cannot be seen from a screen."),
]

Q4 = """Forgeron. We build a five-axis CNC machine — a machine that carves a metal part out of a solid block — together with an AI agent that lets you operate it by speaking or writing in plain language.

Today, a workshop in Mali that receives an order for a complex, precise part usually has to turn it down. The machine that could make it costs around $40,000 imported, and running one requires a CNC programmer: a specialist who writes thousands of lines of machine code by hand, and who is rare and expensive here. So the part is imported instead — weeks of waiting, several times the price.

We build the machine in Bamako for around $15,000, and the AI removes the need for the programmer. The operator says what they want, the agent explains exactly what it is about to do and waits for confirmation before anything moves. A workshop owner who has never written a line of code can make the part himself."""

Q5 = """First, the small machine shops and fabrication workshops around us, and the machinists who work in them. Today they turn away precisely the most valuable orders — the complex, precise ones — and that work leaves the country. With Forgeron they can keep it.

Then their customers: mechanics, equipment repairers, mould makers, anyone who needs a spare part that simply does not exist on the local market and currently waits weeks for an import.

And technical schools and fablabs, where students could finally learn on a five-axis machine instead of only reading about one."""

Q6 = """It meant being taken seriously. I built this machine with my own money, while moving from being a student to being an engineer, and for a long time it was just a thing in a workshop that a few people found interesting. Being selected out of 670 applications from 24 countries told me the problem we are working on is real, and not only to me.

It also placed us somewhere unusual. Most of the cohort builds software; we arrive with a physical machine that cuts metal. That is when I understood that our difference is not a handicap — it is the point."""

Q7 = """To stop having a machine that only works in our own workshop. By the end of the programme I want Forgeron installed in real workshops in Bamako, producing real customer parts, with the owners paying for it — not a free demonstration, a paid pilot. That is the step that turns this from a project into a company.

Alongside that, I want the AI agent to speak the language of the workshop. We have started work on Bambara, which is what people actually speak standing at the machine."""

Q8 = ("In our communities, the blacksmith is the one who can make whatever you ask for. We named "
      "our company Forgeron because that is exactly what we are putting back into the workshop — "
      "except now you ask the machine, in your own words.")

# -------------------------------------------------------------- remplissage
doc = Document(SRC)
T = doc.tables


def fill(tbl, text, row=0, col=0):
    cell = tbl.rows[row].cells[col]
    chunks = [c.strip() for c in text.split("\n\n") if c.strip()]
    cell.text = chunks[0]
    for extra in chunks[1:]:
        cell.add_paragraph(extra)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.size = Pt(11)


fill(T[0], Q1)
fill(T[1], Q2)
fill(T[2], Q3)
for i, (name, role) in enumerate(TEAM, start=1):
    T[3].rows[i].cells[1].text = name
    T[3].rows[i].cells[2].text = role
fill(T[4], Q4)
fill(T[5], Q5)
fill(T[6], Q6)
fill(T[7], Q7)
fill(T[8], Q8)
# T[9] = consentement : volontairement laissé vide, à cocher par Lamine.

doc.save(OUT)
print("Questions remplies : 8/8  (+ tableau equipe : 2 membres)")
print("Laisse vide a dessein : question 9 (consentement) — a cocher par Lamine")
print("Fichier :", OUT)
