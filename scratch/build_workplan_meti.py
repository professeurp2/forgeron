#!/usr/bin/env python3
"""Remplit le modèle METI Innovator Workplan pour Forgeron."""
import shutil
from pathlib import Path
import openpyxl

SRC = Path("/tmp/claude-0/-home-user-forgeron/f115b2cf-c02c-5ae4-8f73-96461eacf843/scratchpad/"
           "meti_in/082e4206-METI_Cohort_Innovator_Workplan.xlsx")
OUT = Path("/home/user/forgeron/scratch/Forgeron_METI_Workplan.xlsx")

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)
ws = wb["Workplan"]

# ---------------------------------------------------------------- en-tête
ws["C4"] = "Forgeron"
ws["C5"] = "Mali"
ws["C6"] = "Lamine SACKO"
ws["C7"] = ("Aboubacar Diamouténé — electronics and assembly technician. "
            "Commercial/business role currently open.")
ws["C8"] = "sackolamine994@gmail.com"

# ---------------------------------------------------------------- partie 1
ws["C11"] = ("a 5-axis CNC machine that works only in our own workshop, with 0 customer parts "
             "produced, 0 documented customer interviews and no validated price")
ws["C14"] = ("a paid pilot in 2 Bamako machine shops, each having produced at least 5 real "
             "customer parts on the machine, at a price the shops have agreed to pay")
ws["C17"] = "April 2027"
ws["C18"] = ("2 signed pilot agreements; at least 10 real customer parts produced (5 or more per "
             "shop), each with its drawing, photos and machining log; at least one payment "
             "received from a shop; 10 documented customer interviews; and a unit-economics model "
             "built on a real, quoted bill of materials rather than estimates.")
ws["C20"] = ("We have only one machine. It travels to Addis Ababa in November and must then serve "
             "both pilots one after the other, so a breakdown, a power cut or transport damage "
             "stops the whole plan. The second risk is commercial: a shop may accept the machine "
             "and still refuse to pay for it.")

# ---------------------------------------------------------------- partie 2
# (#, activité, catégorie, besoin, qté, unité, coût unitaire, UniPod, mois, notes)
EST = "Estimate — to be confirmed by the team"
ROWS = [
 (1, "Run 10 structured customer interviews and select 2 pilot shops",
    "Market access / Partnerships", "Introductions to machine shops and industrial SMEs in Bamako",
    10, "introductions", 0, "Partly", "Dec 2026", "Via UniPod Mali and ENI-ABT networks"),
 (1, "Run 10 structured customer interviews and select 2 pilot shops",
    "Travel / Logistics", "Transport within Bamako for interviews and shop visits",
    15, "trips", 8, "No", "Dec 2026", EST),
 (1, "Run 10 structured customer interviews and select 2 pilot shops",
    "Technical expertise / Mentorship", "Mentor on customer discovery and institutional selling",
    4, "sessions", 50, "Partly", "Dec 2026", EST),

 (2, "Make the machine safe to run outside our own workshop",
    "Materials / Consumables",
    "Hardware emergency stop: NC mushroom button, contactor, relay, wiring",
    1, "set", 80, "No", "Nov 2026",
    "Prerequisite before the machine leaves the workshop — a crashed ESP32 must not leave the spindle running"),
 (2, "Make the machine safe to run outside our own workshop",
    "Technical expertise / Mentorship", "Machine-safety review by a qualified engineer",
    2, "sessions", 75, "Partly", "Nov 2026", EST),
 (2, "Make the machine safe to run outside our own workshop",
    "Materials / Consumables", "Spare parts kit: TB6600 drivers, ESP32 boards, belts, limit switches",
    1, "kit", 150, "No", "Dec 2026", EST),
 (2, "Make the machine safe to run outside our own workshop",
    "Space / Facilities", "Workshop bench and storage at UniPod Mali for preparation work",
    5, "months", 0, "Yes", "Nov 2026", "Already available to us"),

 (3, "Transport and install the machine in each pilot shop",
    "Equipment / Hardware", "Custom transport crate with foam padding",
    1, "crate", 150, "No", "Dec 2026", EST),
 (3, "Transport and install the machine in each pilot shop",
    "Travel / Logistics", "Machine transport, round trip, 2 shops",
    4, "trips", 50, "No", "Jan 2027", EST),
 (3, "Transport and install the machine in each pilot shop",
    "Space / Facilities", "Floor space and mains power inside the pilot shop",
    5, "months", 0, "No", "Jan 2027", "Provided by the pilot shop as part of the agreement"),

 (4, "Run the AI agent fully offline inside the pilot shop",
    "Equipment / Hardware", "Laptop with an 8 GB VRAM GPU to run the local model offline",
    1, "unit", 1200, "No", "Dec 2026",
    "Largest single need. The shops have no reliable internet, so the agent must run locally"),
 (4, "Run the AI agent fully offline inside the pilot shop",
    "Equipment / Hardware", "UPS / backup power for the machine and the PC",
    1, "unit", 250, "No", "Dec 2026", "Power cuts are routine in Bamako — " + EST.lower()),
 (4, "Run the AI agent fully offline inside the pilot shop",
    "Software / Digital tools", "Ollama, FreeCAD, Meshroom (open source)",
    5, "months", 0, "No", "Dec 2026", "Already in use, no licence cost"),

 (5, "Train the shop operator to run the machine in plain language",
    "Skills / Training", "Operator training sessions delivered in the shop",
    4, "sessions", 0, "No", "Jan 2027", "Delivered by our own team"),
 (5, "Train the shop operator to run the machine in plain language",
    "Technical expertise / Mentorship",
    "Native Bambara speakers to review the closed command lexicon and the safety labels",
    3, "sessions", 30, "Partly", "Jan 2027",
    "Shop operators speak Bambara; emergency-stop and alarm labels must be verified by a native speaker"),

 (6, "Produce at least 5 real customer parts in each shop",
    "Materials / Consumables", "Aluminium and steel stock for customer parts and trials",
    60, "kg", 6, "No", "Feb 2027", EST),
 (6, "Produce at least 5 real customer parts in each shop",
    "Materials / Consumables", "Carbide cutting tools and replacements",
    12, "tools", 25, "No", "Feb 2027", EST),
 (6, "Produce at least 5 real customer parts in each shop",
    "Equipment / Hardware", "Metrology set: digital calliper, dial indicator, surface plate",
    1, "set", 250, "Partly", "Feb 2027", "Needed to prove the parts are within tolerance"),
 (6, "Produce at least 5 real customer parts in each shop",
    "Testing / User access", "Access to the pilot shops and to their real customer orders",
    5, "months", 0, "No", "Jan 2027", "This is the core of the pilot"),

 (7, "Validate the price and build the unit-economics model",
    "Technical expertise / Mentorship", "Mentor on pricing and unit economics for hardware ventures",
    4, "sessions", 50, "Partly", "Mar 2027", EST),
 (7, "Validate the price and build the unit-economics model",
    "Other", "Costed bill of materials: supplier quotes for mechanical and electronic parts",
    1, "exercise", 0, "No", "Mar 2027",
    "Our own cost of goods is currently unknown — this is the input the model is missing"),

 (8, "Protect the work and pass investor due diligence",
    "Regulatory / Certification", "Trademark registration for Forgeron (OAPI)",
    1, "filing", 350, "No", "Mar 2027", "unknown — fee to be confirmed with OAPI"),
 (8, "Protect the work and pass investor due diligence",
    "Regulatory / Certification", "Machine-safety conformity assessment for commercial sale",
    1, "assessment", 500, "No", "Apr 2027", "unknown — no quote obtained yet"),

 (9, "Document the evidence for the programme and for investors",
    "Software / Digital tools", "Mobile data for the team during the pilot",
    5, "months", 20, "No", "Jan 2027", EST),
]

r = 27
for row in ROWS:
    for col, val in zip("ABCDEFG", row[:7]):
        ws[f"{col}{r}"] = val
    ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"] = row[7], row[8], row[9]
    r += 1

# ---------------------------------------------------------------- partie 3
ws["B66"] = ("A laptop with an 8 GB VRAM GPU, so the AI agent runs fully offline inside workshops "
             "that have no reliable internet (~USD 1,200).")
ws["B67"] = ("Introductions to machine shops and industrial SMEs in Bamako, plus mentorship on "
             "institutional selling. We have a working machine and no customer pipeline — this is "
             "our weakest point.")
ws["B68"] = ("Material, cutting tools and metrology to produce and verify the pilot parts "
             "(~USD 910), and the safety/certification work needed before we can sell (~USD 850).")

wb.save(OUT)

# ---------------------------------------------------------------- contrôle
total = sum(x[4] * x[6] for x in ROWS)
yes = sum(x[4] * x[6] for x in ROWS if x[7] == "Yes")
partly = sum(x[4] * x[6] for x in ROWS if x[7] == "Partly")
no = sum(x[4] * x[6] for x in ROWS if x[7] == "No")
print(f"Lignes d'activites : {len(ROWS)} (lignes 27 a {r-1}) sur 30 disponibles")
print(f"TOTAL              : {total:>8,} USD")
print(f"  dont Yes         : {yes:>8,} USD")
print(f"  dont Partly      : {partly:>8,} USD")
print(f"  dont No          : {no:>8,} USD")
print(f"ECART DE FINANCEMENT (Partly + No) : {partly+no:,} USD")
print("Fichier :", OUT)
