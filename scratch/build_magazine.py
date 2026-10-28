#!/usr/bin/env python3
"""Insère les réponses Forgeron, en vert, sous chaque question du magazine UniPods."""
from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt

SRC = Path("/tmp/claude-0/-home-user-forgeron/f115b2cf-c02c-5ae4-8f73-96461eacf843/scratchpad/"
           "mag_in/149c560b-UniPods_Magazine.docx")
OUT = Path("/home/user/forgeron/scratch/UniPods_Magazine_Forgeron_LamineSACKO.docx")

GREEN = RGBColor(0x1F, 0x7A, 0x33)
TODO = RGBColor(0xC0, 0x00, 0x00)

A = {}

A["Name, country, UniPod"] = """Lamine SACKO, Mali. I'm a Cohort 1 innovator in the METI UniPods AI Innovation Programme, based at UniPod Mali, hosted by the École Nationale d'Ingénieurs Abderhamane Baba Touré (ENI-ABT) in Bamako. I'm the founder and engineer of Forgeron: I designed the machine and built the control software, the mobile app and the AI agent. Aboubacar Diamouténé is our electronics and assembly technician — he builds and wires the machine, and he finds the faults I can't see from behind a screen. We are two people. The commercial role is still open, and this programme has taught me that it's our biggest gap."""

A["What did you study or do before"] = """I studied mechatronics engineering at the École Nationale d'Ingénieurs Abderhamane Baba Touré (ENI-ABT) in Bamako. But the real beginning came long before that, and it has a name.

In our communities, the forgeron — the blacksmith — is the maker. He is the one who produces whatever the community needs, out of metal and his own hands. I'm not a blacksmith myself, but one lived close to our house, and he made parts for the farmers around us: tools, spares, whatever anyone brought him. Every morning I would go and watch him work. I was a child and I found it extraordinary — not the fire or the noise, but the sheer capability of it. Whatever was asked of him, he could make it.

That never left me. When I finished my studies I didn't want to design machines on paper; I wanted to make things, the way he did. It's also where our name comes from. We called the company Forgeron because that is exactly what we are trying to put back into the workshop: the ability to make whatever is asked for, here."""

A["Describe yourself as an innovator"] = """An engineer who would rather build the machine his country imports than explain why it can't be built here."""

A["4. What problem are you solving"] = """Small machine shops and manufacturing SMEs in West Africa can't produce complex, high-precision parts. A five-axis machine — the kind that can cut a curved surface out of a solid block — costs around $40,000 imported, and running one needs a CNC programmer: a specialist who writes thousands of lines of machine code by hand. Those people are rare in Mali, and expensive. So when a customer walks into a workshop with a complicated part, the owner turns it over in his hand and says he can't make it. The customer imports it instead: three weeks of waiting, three times the price. The people who feel this most are the workshop owners and machinists who lose the work, and their customers — mechanics, equipment repairers, mould makers — who need a part that simply isn't available locally."""

A["5. When did you first notice"] = """[[À COMPLÉTER — c'est la question la plus personnelle du magazine, et celle que les lecteurs retiennent. Raconte UN moment précis : un atelier où tu as vu refuser une pièce, une réparation impossible dans ta famille, un stage. Quand, où, qui était là, ce qui s'est dit. Si tu me donnes la scène en français, même en vrac, je la rédige.]]"""

A["6. Explain your solution"] = """Think of a machine that carves a metal part out of a solid block, following a drawing. Normally, to tell it what to do, someone has to write thousands of lines of coordinates by hand. That's a trained specialist's job, and in Mali those specialists barely exist.

With Forgeron, you simply say what you want, in plain language, the way you'd explain it to a colleague. The machine understands, tells you exactly what it is about to do, and waits for you to say yes before anything moves. If what you asked for would damage the machine or hurt someone, it refuses and explains why.

So a workshop owner who has never written a line of code can run a machine that used to require an engineer."""

A["7. Where does AI actually"] = """The AI understands and proposes. It never calculates. That distinction is the whole safety design.

The AI reads what the operator says and works out which of its 38 machine tools to use — read the position, move an axis, start the spindle, set a work origin — then proposes the action in words. The actual numbers, the toolpaths and the limit checks are computed by ordinary deterministic code, not by the model. Before any movement, the system simulates the entire path; if the tool would hit the table or leave its travel, the command is refused and the machine does not move. The operator confirms every action, and the emergency stop can never be filtered, delayed or overruled by the AI.

We state it as a rule: the AI proposes, deterministic code computes, the operator confirms, and the hardware stops. The AI removes the barrier — it doesn't take the controls."""

A["8. What stage are you at"] = """A working prototype moving into field testing. The machine is built and running in our workshop in Bamako: five axes, three linear and two rotary. The application is around 52,000 lines of code with 342 automated tests, and the AI agent has 38 machine tools. It runs on Gemini when there's internet and on a local model, fully offline, when there isn't — which is the normal case in a Malian workshop.

Right now a user can connect to the machine, see the live position of all five axes, ask for an operation in plain language, read back what the agent proposes, confirm it and watch the machine do it. We have machined a real part this way: a hemispherical dome, 20 mm radius, in 3 hours 35.

What a user cannot do yet is buy one. We have no paying customer. That is the next step."""

A["11. What do you like about the MIT"] = """[[À COMPLÉTER sur MIT — une phrase ou deux : ce que le cours t'a apporté concrètement. Si tu ne l'as pas terminé, dis-le simplement, c'est mieux qu'un compliment vague.]]

From Wadhwani, the concept I now use constantly is describing the customer's problem without ever mentioning your own solution. The first time I tried it I couldn't do it — every sentence came back to my machine. It sounds like a writing exercise, but it changed the way I talk to workshop owners: I ask them about the orders they turned down last month, instead of telling them what my machine can do. The conversations are completely different now."""

A["13. What has changed in your solution"] = """A great deal, and most of it isn't visible from the outside.

When I applied, the AI agent only worked with a cloud model — which is useless in a workshop with no reliable internet. It now runs on a local model, fully offline, and the number of machine tools it can call has gone from a handful to 38. We built a complete path from a 3D drawing to a finished toolpath, including continuous five-axis motion over free-form surfaces. We also added something I hadn't planned: you can photograph a broken part from several angles and get a 3D model back, so a workshop can reproduce a part it has no drawing for. The interface now speaks 16 languages, and we have started on Bambara — the language actually spoken on the shop floor.

And we learned to measure. We built a benchmark that replays fixed scenarios, so when we say the agent has improved, we can prove it instead of claiming it."""

A["14. What was your hardest moment"] = """The machine has broken down three times in ways that took days to find — not a loose wire, but faults where the symptom was nowhere near the cause. Each time I had to stop building features and go back to the hardware, with deadlines getting closer.

What got me through it was writing down every failure: what triggered it, how we detected it, what guard we added afterwards. That file is now part of the project. Later, when the AI model started misbehaving — inventing machine code, drifting into another language mid-sentence — we added those to the same list. The failures became the safety rules.

I don't think I would trust the machine today if it had never broken.

[[Si tu préfères un moment plus personnel — le financement, la fatigue, un moment de doute — dis-le-moi et je réécris. Celui-ci est vrai et documenté, mais c'est toi qui choisis ce que tu racontes.]]"""

A["18. How many people have used"] = """None outside our own team yet, and I would rather say that plainly than inflate it. The machine works; it has not yet been in a customer's hands.

What we do have is measured. The machine has produced a real part end to end. The application passes 342 automated tests. And we have measured the AI agent itself: on our offline local model it calls the correct tool 28 to 29 times out of 30 in normal use, and we know exactly where that degrades, because we measured that too.

We have had a first conversation with a Bamako workshop, Kouratechnique, about testing the machine in real production. Our goal for the next six months is a paid pilot in two workshops, with real customer parts coming off the machine."""

A["19. Share one story"] = """I don't have one yet, and I won't invent one.

The closest thing I have is my own. The first time the machine made a part from a sentence, I asked for a dome in plain language, read back what the agent proposed, confirmed it, and three and a half hours later I was holding it. What changed that day is that the project stopped being an argument about whether it was possible.

[[Facultatif — une phrase sur ce que tu as ressenti ce jour-là, ou sur la réaction d'Aboubacar. Une ligne suffit.]]"""

A["20. Where do you want this"] = """In 12 months: machines in a handful of Bamako workshops, paid for, with owners taking on precision orders they used to turn away. By then I want to know our real cost of production and have a price that customers have actually agreed to pay — not an estimate.

In 3 years: Forgeron machines built in Mali and sold across West Africa, and an AI operator that speaks the languages of the workshop, starting with Bambara, so the person running the machine doesn't need French or English to use it."""

A["21. What do you need most"] = """Customers, first. We have a working machine and no customer pipeline. Introductions to workshops and industrial SMEs willing to run a pilot would be worth more to us right now than money.

Second, a commercial co-founder. I can build the machine; I can't yet sell it, and I've stopped pretending otherwise.

Third, funding to build the first production batch, and compute to train speech models for Bambara, where almost no data exists today."""

A["22. What would a partnership"] = """Japan leads the world in machine tools and industrial robotics. Forgeron is a machine-tool company — in a cohort where most ventures are software, I think that makes us an unusual but natural fit.

Concretely, a Japanese manufacturer could help us with exactly what we are weakest at: turning one hand-built machine into something reproducible. Precision at scale, quality control, component sourcing, and the discipline of making the same thing twice. That is a century of knowledge we don't have.

And it wouldn't be one-way. We are learning how to build a machine that survives in a workshop with unreliable power, no internet and no trained programmer. That problem exists in a great many more places than Mali."""

A["23. What advice would you give"] = """Build the smallest ugly version, and build it now. Not the business plan, not the presentation — the thing itself. My machine was made from parts I could find in Bamako and it looked bad for a long time, but it moved. The day something moves, people stop arguing with you about whether it's possible and start asking what it costs.

And don't wait until you feel ready. I still don't."""

A["24. Finish this sentence"] = """"Before UniPods I thought the hard part was building the machine; now I know the hard part is finding the person who will pay for it — and I'm finally learning how to do that.\""""

A["25. One line you"] = """We built a five-axis machine in Bamako that you operate by talking to it. If that's possible here, with what we had, then a great deal of what Africa imports could be made here instead."""

# ------------------------------------------------------------------ reconstruction
# Les sections du modèle regroupent plusieurs questions dans un seul paragraphe,
# séparées par des sauts de ligne. On reconstruit : une ligne = un paragraphe.

SECTIONS = {"About you", "2. The problem and the solution",
            "3. Your journey with the programme and UniPods", "5. Impact and evidence",
            "6. Looking ahead", "7. Reflections"}

src = Document(SRC)
lines = []
for p in src.paragraphs:
    for ln in p.text.split("\n"):
        ln = ln.strip()
        if ln:
            lines.append(ln)

doc = Document()


def add(text, *, color=None, bold=False, italic=False, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(11)
    if color is not None:
        r.font.color.rgb = color
    return p


placed, missing = [], set(A)
for ln in lines:
    if ln in SECTIONS:
        add(ln, bold=True, space_before=14)
        continue
    if ln.startswith("Please just simply put"):
        add(ln, italic=True)
        continue

    add(ln, bold=True, space_before=10)

    # la question correspond-elle à une réponse ?
    for key, answer in A.items():
        probe = key.split(". ", 1)[-1] if key[0].isdigit() else key
        if probe.lower() in ln.lower():
            for chunk in answer.split("\n\n"):
                chunk = chunk.strip()
                todo = chunk.startswith("[[")
                add(chunk, color=TODO if todo else GREEN, italic=todo)
            placed.append(key)
            missing.discard(key)
            break

doc.save(OUT)
print(f"Questions du modèle : {len([l for l in lines if l not in SECTIONS])}")
print(f"Réponses placées    : {len(placed)}/{len(A)}")
if missing:
    print("NON TROUVÉES :", sorted(missing))
print(f"Blocs [[À COMPLÉTER]] : {sum(1 for v in A.values() if '[[' in v)}")
print("Fichier :", OUT)
