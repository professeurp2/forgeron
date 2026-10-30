#!/usr/bin/env python3
"""Remplit le formulaire de profil pour l'exposition numérique du PAYF'26."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

SRC = Path("/tmp/claude-0/-home-user-forgeron/f115b2cf-c02c-5ae4-8f73-96461eacf843/scratchpad/"
           "payf_in/44fa4ef3-PAYF26_Digital_Exhibition_Profile_Form.docx")
OUT = Path("/home/user/forgeron/scratch/PAYF26_Digital_Profile_Forgeron_LamineSACKO.docx")

DESC = ("Forgeron builds an affordable 5-axis CNC machine in Bamako, Mali, operated in plain "
        "language by an AI agent. Workshops that once turned away complex precision parts — and "
        "imported them — can now make them locally, without a specialist programmer.")

ANSWERS = {
    "1. Full name": "Lamine SACKO",
    "2. Country": "Mali",
    "3. Business / organisation / innovation name": "Forgeron",
    "4. What does your work do?": DESC,
    "5. Website or link of your work": "https://github.com/professeurp2/forgeron",
    "6. Your Social Media Links": "",
}

doc = Document(SRC)
tbl = doc.tables[0]

placed = []
for i, row in enumerate(tbl.rows):
    label = row.cells[0].text.strip()
    for key, val in ANSWERS.items():
        if label.startswith(key.split("(")[0].strip()) and val:
            # la valeur se trouve dans le paragraphe placeholder sous le libellé
            cell = row.cells[0]
            target = None
            for p in cell.paragraphs:
                if p.text.strip().startswith("["):
                    target = p
                    break
            if target is None:
                target = cell.add_paragraph()
            for r in list(target.runs):
                r.text = ""
            run = target.add_run(val) if not target.runs else target.runs[0]
            run.text = val
            run.font.size = Pt(11)
            run.italic = False
            placed.append(key)
            break

doc.save(OUT)
print("Champs remplis :", len(placed))
for k in placed:
    print("  -", k)
print("Mots dans la description :", len(DESC.replace("—", " ").split()))
print("Laissé vide : réseaux sociaux (facultatif) · case d'autorisation (à cocher par Lamine)")
print("Fichier :", OUT)
